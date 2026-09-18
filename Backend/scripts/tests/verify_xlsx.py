"""Check the rewritten workbook against the code, not against my intent.

Every spelling the 'Size Prefixes' sheet promises is fed to the real
_friendly_size; every tag word the 'Size Tag Words' sheet promises is looked up
in the real SIZE_ALIAS_GROUPS. A doc that claims something the script does not do
is worse than no doc - it is what sends a designer off renaming panels.
"""
import copy
import json
import os
import re
import subprocess
import zipfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))   # Backend/scripts/tests -> repo root
PUB = os.path.join(ROOT, "Frontend", "my-app", "public", "Pattern_Naming_Reference.xlsx")
DOCS = os.path.join(ROOT, "Backend", "scripts", "docs", "Pattern_Naming_Reference.xlsx")
# The pre-change copy, if one was kept. Only the styling comparison needs it;
# everything that matters - does the sheet promise what the code does - runs
# without it. Put a copy here to re-enable that section after an edit.
OLD = os.path.join(HERE, "xlsx_backup", "public_Pattern_Naming_Reference.xlsx")
HAVE_OLD = os.path.exists(OLD)

ia = {"re": re}
src = open(os.path.join(ROOT, "Backend", "services", "illustrator_automation.py"), encoding="utf-8").read()
exec(src[src.index("_FRIENDLY_SIZE_MAP = {"):src.index("def pattern_files_needed")], ia)
friendly, youth_file = ia["_friendly_size"], ia["is_youth_pattern_size"]

fails, checks = [], 0


def check(name, got, want):
    global checks
    checks += 1
    if got != want:
        fails.append(f"{name}\n      got  {got!r}\n      want {want!r}")


# ------------------------------------------------- the files are still valid
for p in (PUB, DOCS):
    check(f"{os.path.basename(os.path.dirname(p))} zip intact", zipfile.ZipFile(p).testzip(), None)
wb, wb2 = (openpyxl.load_workbook(p) for p in (PUB, DOCS))
old = openpyxl.load_workbook(OLD) if HAVE_OLD else None
if old:
    check("sheet names unchanged", wb.sheetnames, old.sheetnames)
check("both copies identical", [[[c.value for c in r] for r in wb[s].iter_rows()] for s in wb.sheetnames],
      [[[c.value for c in r] for r in wb2[s].iter_rows()] for s in wb2.sheetnames])

# ------------------------------------------------------ styling was preserved
# The lesson from PHR 197: a data row must never wear the header's style.
for sheet in (("Size Prefixes", "Size Tag Words", "Read Me", "Pattern Names") if old else ()):
    ws, ows = wb[sheet], old[sheet]
    hdr_row = 2 if sheet == "Pattern Names" else 1
    for col in range(1, ws.max_column + 1):
        h = ws.cell(row=hdr_row, column=col)
        o = ows.cell(row=hdr_row, column=col)
        check(f"{sheet} header col{col} font", (h.font.b, h.font.color and h.font.color.rgb),
              (o.font.b, o.font.color and o.font.color.rgb))
    sample = 3 if sheet == "Pattern Names" else 2
    for r in range(sample, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is None:
            continue
        for col in range(1, ws.max_column + 1):
            c, ref = ws.cell(row=r, column=col), ows.cell(row=sample, column=col)
            check(f"{sheet} r{r}c{col} not wearing the header style",
                  (c.font.b, c.border.bottom.style), (ref.font.b, ref.border.bottom.style))
    check(f"{sheet} column widths", {k: v.width for k, v in ws.column_dimensions.items()},
          {k: v.width for k, v in ows.column_dimensions.items()})
    check(f"{sheet} freeze pane", ws.freeze_panes, ows.freeze_panes)

# ------------------------------- what the sheet PROMISES is what the code DOES
ws = wb["Size Prefixes"]
promised = []
for r in range(2, ws.max_row + 1):
    order_says, canonical, example, kind, others = (ws.cell(row=r, column=c).value for c in range(1, 6))
    if not order_says or canonical in ("used exactly as written", "the same as without the A"):
        continue
    spellings = [s.strip() for s in re.split(r"\u00b7|\bor\b", str(order_says)) if s.strip()]
    spellings += [s.strip() for s in str(others or "").split("\u00b7") if s.strip()]
    for s in spellings:
        # Prose, not a spelling - the two rows that explain rather than list.
        if len(s.split()) > 4 or s in ("\u2014", "the same code"):
            continue
        if "to" in s.split() or s.endswith("..."):
            continue
        promised.append((s, canonical, kind, r))

by_row = {}
for spelling, canonical, kind, r in promised:
    got = friendly(spelling)
    want = canonical if canonical != "the same code" else got
    if str(canonical).startswith("W-") or canonical in ("XS", "Small", "Medium", "Large", "XL",
                                                        "2XL", "3XL", "4XL", "YXS", "YS", "YM", "YL", "YXL"):
        by_row.setdefault((r, canonical), []).append((spelling, got))
for (r, canonical), pairs in sorted(by_row.items()):
    wrong = [(s, g) for s, g in pairs if g != canonical]
    check(f"Size Prefixes row {r} ({canonical}) - every promised spelling resolves", wrong, [])

# The 'Type' column claims which pattern file the size is cut from.
for r in range(2, ws.max_row + 1):
    canonical, kind = ws.cell(row=r, column=2).value, ws.cell(row=r, column=4).value
    if not str(canonical or "").startswith("W-"):
        continue
    check(f"Size Prefixes row {r} ({canonical}) file column",
          youth_file(canonical), "youth" in str(kind))
    check(f"Size Prefixes row {r} ({canonical}) example name",
          ws.cell(row=r, column=3).value, canonical + " Front")

# --------------------------- the tag words the sheet promises really do rename
r = subprocess.run(["node", os.path.join(HERE, "tag_groups.js")], capture_output=True, text=True, cwd=HERE)
groups = json.loads(r.stdout) if r.returncode == 0 else None
check("SIZE_ALIAS_GROUPS readable from the JSX", groups is not None, True)
if groups:
    norm = lambda s: re.sub(r"[^a-z0-9]", "", str(s).lower())
    ws = wb["Size Tag Words"]
    for row in range(2, ws.max_row + 1):
        size, anysize, scoped = (ws.cell(row=row, column=c).value for c in (1, 2, 3))
        if not size or not str(scoped) or str(scoped) == "\u2014":
            continue
        want_key = norm(size)
        group = next((g for g in groups if want_key in g), None)
        check(f"Size Tag Words r{row}: '{size}' has a scoped group", group is not None, True)
        if group:
            missing = [w.strip() for w in str(scoped).split("\u00b7")
                       if w.strip() and norm(w) not in group]
            check(f"Size Tag Words r{row}: '{size}' promises only words the group holds", missing, [])

print(f"{checks - len(fails)}/{checks} passed")
for f in fails:
    print("  FAIL  " + f)
raise SystemExit(1 if fails else 0)
